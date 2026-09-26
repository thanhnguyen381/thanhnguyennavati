"""
Tạo 3 file CV mẫu để test hệ thống chấm điểm
"""
import docx
import os

output_dir = r"e:\THANH NGUYEN\TN-Vibecoding\TN-APP CHAM DIEM CV TU DONG CHO HR\test_cvs"
os.makedirs(output_dir, exist_ok=True)

# =============================================
# CV 1: Ứng viên XUẤT SẮC — Marketing
# =============================================
doc1 = docx.Document()
doc1.add_heading("TRẦN THỊ MINH ANH", 0)
doc1.add_paragraph("📧 minhhanh@email.com | 📞 0901 234 567 | TP.HCM")
doc1.add_paragraph("LinkedIn: linkedin.com/in/minhhanh | Portfolio: minhhanh.design")

doc1.add_heading("MỤC TIÊU NGHỀ NGHIỆP", level=1)
doc1.add_paragraph("Marketing Executive với 4 năm kinh nghiệm thực chiến trong Digital Marketing, mong muốn đóng góp vào sự phát triển bền vững của công ty.")

doc1.add_heading("KINH NGHIỆM LÀM VIỆC", level=1)
doc1.add_paragraph("2022 – Nay | Senior Marketing Executive | Brand ABC Corp")
doc1.add_paragraph("- Quản lý ngân sách quảng cáo Facebook Ads & Google Ads lên đến 500 triệu/tháng")
doc1.add_paragraph("- Tăng tỷ lệ chuyển đổi (conversion rate) từ 1.2% lên 3.8% sau 6 tháng")
doc1.add_paragraph("- Xây dựng chiến lược content cho fanpage 200K followers")
doc1.add_paragraph("- Thiết kế creative bằng Canva, Photoshop và quản lý đội ngũ 3 người")
doc1.add_paragraph("2020 – 2022 | Marketing Executive | StartupXYZ")
doc1.add_paragraph("- Chạy chiến dịch Google Ads, đạt ROAS 4.5x")
doc1.add_paragraph("- Viết content SEO, tăng traffic organic 200%")

doc1.add_heading("KỸ NĂNG", level=1)
doc1.add_paragraph("Facebook Ads, Google Ads, TikTok Ads, Canva, Photoshop, SEO, Content Writing, Google Analytics, Email Marketing, A/B Testing")

doc1.add_heading("HỌC VẤN", level=1)
doc1.add_paragraph("Cử nhân Marketing — Đại học Kinh tế TP.HCM (2016–2020) — GPA: 3.5/4.0")
doc1.add_paragraph("Chứng chỉ: Google Ads Certified, Facebook Blueprint, HubSpot Inbound Marketing")

doc1.save(os.path.join(output_dir, "CV1_TranThiMinhAnh_XUAT_SAC.docx"))
print("✅ Tạo xong CV 1: Trần Thị Minh Anh (Xuất sắc)")

# =============================================
# CV 2: Ứng viên TRUNG BÌNH — Marketing
# =============================================
doc2 = docx.Document()
doc2.add_heading("NGUYỄN VĂN BÌNH", 0)
doc2.add_paragraph("Email: binh.nguyen@gmail.com | SĐT: 0912 555 888 | Hà Nội")

doc2.add_heading("GIỚI THIỆU BẢN THÂN", level=1)
doc2.add_paragraph("Tôi là sinh viên mới tốt nghiệp ngành Truyền thông, đam mê Marketing và muốn học hỏi thêm kinh nghiệm.")

doc2.add_heading("KINH NGHIỆM", level=1)
doc2.add_paragraph("2023 – 2024 | Thực tập sinh Marketing | Công ty DEF")
doc2.add_paragraph("- Hỗ trợ đăng bài lên fanpage Facebook")
doc2.add_paragraph("- Viết caption và thiết kế ảnh bằng Canva")
doc2.add_paragraph("- Báo cáo số liệu hàng tuần")
doc2.add_paragraph("Freelance (6 tháng) | Chạy Facebook Ads cho shop thời trang nhỏ")
doc2.add_paragraph("- Ngân sách 5-10 triệu/tháng, đạt được một số đơn hàng")

doc2.add_heading("KỸ NĂNG", level=1)
doc2.add_paragraph("Canva, Facebook cơ bản, Word, Excel, viết content")

doc2.add_heading("HỌC VẤN", level=1)
doc2.add_paragraph("Cử nhân Truyền thông — Đại học Khoa học Xã hội và Nhân văn (2020–2024)")

doc2.save(os.path.join(output_dir, "CV2_NguyenVanBinh_TRUNG_BINH.docx"))
print("✅ Tạo xong CV 2: Nguyễn Văn Bình (Trung bình)")

# =============================================
# CV 3: Ứng viên KHÔNG PHÙ HỢP — IT dev
# =============================================
doc3 = docx.Document()
doc3.add_heading("LÊ HOÀNG PHÚC", 0)
doc3.add_paragraph("Email: phuc.le@dev.com | SĐT: 0933 777 111 | Đà Nẵng")

doc3.add_heading("MỤC TIÊU", level=1)
doc3.add_paragraph("Lập trình viên Backend với 3 năm kinh nghiệm, tìm kiếm cơ hội phát triển trong lĩnh vực phần mềm.")

doc3.add_heading("KINH NGHIỆM", level=1)
doc3.add_paragraph("2021 – Nay | Backend Developer | Tech Solutions Co.")
doc3.add_paragraph("- Phát triển REST API bằng Python (Django, FastAPI)")
doc3.add_paragraph("- Quản lý database PostgreSQL, Redis")
doc3.add_paragraph("- Triển khai CI/CD với Docker, GitHub Actions")
doc3.add_paragraph("2019 – 2021 | Junior Developer | Freelance")
doc3.add_paragraph("- Xây dựng website thương mại điện tử bằng PHP/Laravel")

doc3.add_heading("KỸ NĂNG", level=1)
doc3.add_paragraph("Python, Django, FastAPI, PostgreSQL, Docker, Git, Linux, REST API, Redis")

doc3.add_heading("HỌC VẤN", level=1)
doc3.add_paragraph("Kỹ sư Công nghệ Thông tin — Đại học Bách Khoa Đà Nẵng (2015–2019)")
doc3.add_paragraph("Chứng chỉ: AWS Certified Developer, Docker Certified Associate")

doc3.save(os.path.join(output_dir, "CV3_LeHoangPhuc_KHONG_PHU_HOP.docx"))
print("✅ Tạo xong CV 3: Lê Hoàng Phúc (Không phù hợp - IT Dev)")

print(f"\n📁 Đã lưu tất cả CV vào: {output_dir}")
print("\nDùng JD sau để test:")
print("""
=== JD TEST ===
Vị trí: Marketing Executive
Yêu cầu:
- 2+ năm kinh nghiệm Marketing, Digital Marketing
- Thành thạo Facebook Ads, Google Ads
- Biết dùng Canva, Photoshop cơ bản
- Kỹ năng viết content tốt
- Tốt nghiệp Đại học chuyên ngành Marketing hoặc liên quan
""")
