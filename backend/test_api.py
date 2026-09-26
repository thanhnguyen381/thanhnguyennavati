"""
Script test nhanh API chấm điểm CV
Chạy lệnh: ..\venv\Scripts\python test_api.py
"""
import requests
import io
import docx
import tempfile
import os

API_URL = "http://localhost:8000/score"

# ---- 1. Tạo CV giả dạng DOCX ----
def create_sample_cv() -> bytes:
    doc = docx.Document()
    doc.add_heading("NGUYỄN VĂN AN", 0)
    doc.add_paragraph("Email: an.nguyen@email.com | SĐT: 0912 345 678 | TP.HCM")
    doc.add_heading("MỤC TIÊU NGHỀ NGHIỆP", level=1)
    doc.add_paragraph("Frontend Developer với 4 năm kinh nghiệm, mong muốn đóng góp vào các dự án công nghệ lớn.")
    doc.add_heading("KINH NGHIỆM LÀM VIỆC", level=1)
    doc.add_paragraph("2021 – nay: Senior Frontend Developer tại Tech Corp")
    doc.add_paragraph("- Xây dựng và tối ưu giao diện người dùng bằng React.js, TypeScript, Next.js")
    doc.add_paragraph("- Tích hợp REST API, GraphQL")
    doc.add_paragraph("- Mentor cho 3 junior developer")
    doc.add_paragraph("2019 – 2021: Frontend Developer tại StartupXYZ")
    doc.add_paragraph("- Phát triển landing page với HTML/CSS/JavaScript")
    doc.add_heading("KỸ NĂNG", level=1)
    doc.add_paragraph("React.js, TypeScript, Next.js, Tailwind CSS, Git, Figma, REST API, Redux")
    doc.add_heading("HỌC VẤN", level=1)
    doc.add_paragraph("Cử nhân Công nghệ Thông tin — Đại học Bách Khoa TP.HCM (2015–2019)")
    doc.add_paragraph("Chứng chỉ: AWS Cloud Practitioner, Google UX Design")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

# ---- 2. JD mẫu ----
JD_TEXT = """
Vị trí: Senior Frontend Developer

Yêu cầu:
- Tối thiểu 3 năm kinh nghiệm với React.js và TypeScript
- Có kinh nghiệm với Next.js, Tailwind CSS
- Thành thạo Git, quy trình CI/CD
- Hiểu biết về UX/UI, có thể làm việc với Figma
- Kỹ năng giao tiếp tốt, có khả năng mentor junior

Ưu tiên:
- Kinh nghiệm với GraphQL, Redux
- Có chứng chỉ AWS hoặc GCP
- Từng làm việc tại môi trường Agile/Scrum

Học vấn: Tốt nghiệp Đại học chuyên ngành CNTT hoặc liên quan
"""

# ---- 3. Gửi request ----
print("=" * 55)
print("  🚀 TEST API — AI CV Scorer")
print("=" * 55)
print(f"📡 Gửi request đến: {API_URL}")
print("📄 CV: Nguyễn Văn An (Senior Frontend Developer)")
print("-" * 55)

cv_bytes = create_sample_cv()

# Tạo file tạm thời để gửi
with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
    tmp.write(cv_bytes)
    tmp_path = tmp.name

try:
    with open(tmp_path, "rb") as f:
        files = {"cv": ("nguyen_van_an_cv.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        data  = {"jd_text": JD_TEXT}
        resp  = requests.post(API_URL, files=files, data=data, timeout=60)

    print(f"📶 HTTP Status: {resp.status_code}")

    result = resp.json()
    if result.get("success"):
        d = result["data"]
        print("\n✅ CHẤM ĐIỂM THÀNH CÔNG!\n")
        print(f"  🎯 Tổng điểm    : {d.get('tong_diem')}/10")
        print(f"  🛠️  Kỹ năng      : {d.get('ky_nang', {}).get('diem')}/10 — {d.get('ky_nang', {}).get('nhan_xet', '')[:80]}")
        print(f"  💼 Kinh nghiệm  : {d.get('kinh_nghiem', {}).get('diem')}/10 — {d.get('kinh_nghiem', {}).get('nhan_xet', '')[:80]}")
        print(f"  🎓 Học vấn      : {d.get('hoc_van', {}).get('diem')}/10 — {d.get('hoc_van', {}).get('nhan_xet', '')[:80]}")
        print(f"\n  💪 Điểm mạnh   : {d.get('diem_manh', '')}")
        print(f"  ⚠️  Điểm yếu    : {d.get('diem_yeu', '')}")
        print(f"\n  📌 Khuyến nghị : {d.get('khuyen_nghi', '')}")
    else:
        print(f"\n❌ LỖI: {result.get('error')}")
except Exception as e:
    print(f"\n❌ Không thể kết nối: {e}")
finally:
    os.unlink(tmp_path)

print("\n" + "=" * 55)
