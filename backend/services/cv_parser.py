import fitz  # PyMuPDF
import docx
import io

def extract_text_from_pdf(file_stream: bytes) -> str:
    """Trích xuất văn bản từ file PDF."""
    text = ""
    try:
        doc = fitz.open(stream=file_stream, filetype="pdf")
        for page in doc:
            text += page.get_text() + "\n"
        doc.close()
    except Exception as e:
        print(f"Lỗi khi đọc PDF: {e}")
    return text.strip()

def extract_text_from_docx(file_stream: bytes) -> str:
    """Trích xuất văn bản từ file DOCX (Word), bao gồm cả đoạn văn và bảng biểu."""
    text = ""
    try:
        doc = docx.Document(io.BytesIO(file_stream))
        # Lấy text từ đoạn văn (paragraphs)
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text.strip() + "\n"
                
        # Lấy text từ bảng biểu (tables) - rất quan trọng vì nhiều CV dùng bảng
        for table in doc.tables:
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_data.append(cell.text.strip())
                if row_data:
                    text += " | ".join(row_data) + "\n"
                    
    except Exception as e:
        print(f"Lỗi khi đọc DOCX: {e}")
    return text.strip()

def extract_text(file_stream: bytes, filename: str) -> str:
    """Xác định loại file và gọi hàm trích xuất tương ứng."""
    filename = filename.lower()
    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file_stream)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(file_stream)
    elif filename.endswith(".doc"):
        return "Hệ thống hiện tại chỉ hỗ trợ định dạng .docx cho Word. Vui lòng chuyển đổi .doc sang .docx"
    else:
        return "Định dạng không được hỗ trợ. Vui lòng tải lên file PDF hoặc DOCX."
